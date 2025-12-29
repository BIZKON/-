import fitz  # PyMuPDF
from docx import Document as DocxDocument
from typing import List, Tuple
import io
import tiktoken
import logging

logger = logging.getLogger(__name__)


def count_tokens(text: str, model: str = "gpt-4") -> int:
    """Count tokens in text using tiktoken."""
    try:
        encoding = tiktoken.encoding_for_model(model)
        return len(encoding.encode(text))
    except Exception:
        # Fallback to approximate count
        return len(text) // 4


def chunk_text(text: str, max_tokens: int = 500, overlap: int = 50) -> List[str]:
    """Split text into chunks with overlap."""
    chunks = []
    sentences = text.replace('\n', ' ').split('. ')
    current_chunk = []
    current_tokens = 0

    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue

        sentence_tokens = count_tokens(sentence)

        if current_tokens + sentence_tokens > max_tokens and current_chunk:
            chunk_text = '. '.join(current_chunk) + '.'
            chunks.append(chunk_text)

            # Keep overlap
            overlap_sentences = []
            overlap_tokens = 0
            for s in reversed(current_chunk):
                s_tokens = count_tokens(s)
                if overlap_tokens + s_tokens <= overlap:
                    overlap_sentences.insert(0, s)
                    overlap_tokens += s_tokens
                else:
                    break

            current_chunk = overlap_sentences
            current_tokens = overlap_tokens

        current_chunk.append(sentence)
        current_tokens += sentence_tokens

    if current_chunk:
        chunks.append('. '.join(current_chunk) + '.')

    return chunks


def extract_text_from_pdf(file_data: bytes) -> str:
    """Extract text from PDF file."""
    try:
        doc = fitz.open(stream=file_data, filetype="pdf")
        text_parts = []

        for page in doc:
            text_parts.append(page.get_text())

        doc.close()
        return '\n'.join(text_parts)
    except Exception as e:
        logger.error(f"Error extracting PDF text: {e}")
        raise


def extract_text_from_docx(file_data: bytes) -> str:
    """Extract text from DOCX file."""
    try:
        doc = DocxDocument(io.BytesIO(file_data))
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text)
                if row_text:
                    text_parts.append(' | '.join(row_text))

        return '\n'.join(text_parts)
    except Exception as e:
        logger.error(f"Error extracting DOCX text: {e}")
        raise


def process_document(file_data: bytes, filename: str) -> Tuple[str, List[str]]:
    """Process document and return text and chunks."""
    filename_lower = filename.lower()

    if filename_lower.endswith('.pdf'):
        text = extract_text_from_pdf(file_data)
    elif filename_lower.endswith('.docx'):
        text = extract_text_from_docx(file_data)
    elif filename_lower.endswith('.txt'):
        text = file_data.decode('utf-8')
    else:
        raise ValueError(f"Unsupported file type: {filename}")

    chunks = chunk_text(text)
    return text, chunks
